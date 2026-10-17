"""Petgrave.io-branded Markdown -> Word for compiled Blueprint strategies (Python port of md2docx.js).

    pip install python-docx
    python docs/data-strategy/compile_docx.py editions/petgrave-idl/strategy.md out.docx

Supports the strategy template's subset: front matter, # / ## headings, paragraphs,
- bullets, pipe tables, > callouts, ![caption](figure.png) figures (numbered, kept
with their caption), **bold** and *italic*.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY, SLATE, STEEL, ASH, MIST, CLOUD, TEXT, GREY = "1E2A44", "4A5878", "7E8AA3", "AEB6C6", "D5DAE3", "EEF0F3", "161B26", "5B6272"
FONT = "Calibri"


def rgb(h):
    return RGBColor.from_string(h)


def runs(par, s, size=10.5, color=TEXT, bold=False, italic=False):
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", s):
        if not tok:
            continue
        b, i, t = bold, italic, tok
        if tok.startswith("**"):
            b, t = True, tok[2:-2]
        elif tok.startswith("*"):
            i, t = True, tok[1:-1]
        elif tok.startswith("`"):
            t = tok[1:-1]
        r = par.add_run(t)
        r.font.name, r.font.size, r.font.bold, r.font.italic, r.font.color.rgb = FONT, Pt(size), b, i, rgb(color)


def shade(cell, hex_):
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_)
    tc.append(sh)


def borders(table, color=ASH, size=4):
    tbl = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:color"), color)
        b.append(e)
    tbl.append(b)


def keep_next(par):
    par.paragraph_format.keep_with_next = True


def main(src, out):
    src = Path(src)
    md = src.read_text(encoding="utf-8").replace("\r\n", "\n")
    meta = {}
    fm = re.match(r"^---\n(.*?)\n---\n", md, re.S)
    if fm:
        for line in fm.group(1).splitlines():
            k, _, v = line.partition(":")
            meta[k.strip()] = v.strip()
        md = md[fm.end():]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    hp = sec.header.paragraphs[0]
    runs(hp, meta.get("header", ""), size=8, color=GREY)
    content_w = Cm(21.0 - 4.4)

    # Title block
    p = doc.add_paragraph(); runs(p, meta.get("title", ""), size=24, color=NAVY, bold=True)
    if meta.get("subtitle"):
        p = doc.add_paragraph(); runs(p, meta["subtitle"], size=13, color=SLATE)
    p = doc.add_paragraph()
    runs(p, " · ".join(x for x in (meta.get("edition"), meta.get("date")) if x), size=9.5, color=GREY)
    if meta.get("status"):
        p = doc.add_paragraph(); runs(p, "Status: " + meta["status"], size=9.5, color=GREY, italic=True)

    lines = md.split("\n")
    fig = 0
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue
        m = re.match(r"^(#{1,3}) (.*)", line)
        if m:
            lvl = len(m.group(1))
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            runs(p, m.group(2), size={1: 15, 2: 12.5, 3: 11}[lvl], color=NAVY, bold=True); keep_next(p)
            i += 1; continue
        m = re.match(r"^!\[(.*)\]\((.*)\)", line)
        if m:
            fig += 1
            img = src.parent / m.group(2)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; keep_next(p)
            p.add_run().add_picture(str(img), width=content_w)
            c = doc.add_paragraph()
            runs(c, f"**Figure {fig}.** " + m.group(1), size=9, color=GREY)
            c.paragraph_format.space_after = Pt(10)
            i += 1; continue
        if line.startswith(">"):
            block = []
            while i < len(lines) and lines[i].startswith(">"):
                block.append(lines[i].lstrip("> ").rstrip()); i += 1
            t = doc.add_table(rows=1, cols=1); borders(t, ASH, 6); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = t.rows[0].cells[0]; shade(cell, CLOUD)
            cell.paragraphs[0].text = ""
            for k, b in enumerate(x for x in block if x):
                cp = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                runs(cp, b, size=10.5, color=NAVY)
            doc.add_paragraph()
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells):
                    rows.append(cells)
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); borders(t)
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.rows[r].cells[c]; cell.paragraphs[0].text = ""
                    if r == 0:
                        shade(cell, NAVY); runs(cell.paragraphs[0], val, size=9, color="FFFFFF", bold=True)
                    else:
                        runs(cell.paragraphs[0], val, size=9)
            doc.add_paragraph()
            continue
        m = re.match(r"^(\s*)[-*] (.*)", line)
        if m:
            p = doc.add_paragraph(style="List Bullet"); runs(p, m.group(2))
            i += 1; continue
        m = re.match(r"^\d+\. (.*)", line)
        if m:
            p = doc.add_paragraph(style="List Number"); runs(p, m.group(1))
            i += 1; continue
        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|>|\||!\[|\s*[-*] |\d+\. )", lines[i]):
            para.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(); runs(p, " ".join(para)); p.paragraph_format.space_after = Pt(6)
    doc.save(out)
    print("wrote", out, "with", fig, "figures")


if __name__ == "__main__":
    main(*sys.argv[1:3])
