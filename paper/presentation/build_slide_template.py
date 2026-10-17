"""Build the NASEF 2026 slide template from NASEF_SLIDE_PLAN.md.

    pip install python-pptx
    python paper/presentation/build_slide_template.py

Writes paper/presentation/NASEF2026_IDL_slide_template.pptx. Dashed boxes are
placeholders that name the figure or content to drop in; solid elements are final.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FIG_IDL = ROOT / "docs/data-strategy/editions/petgrave-idl/figures"
FIG_OWN = ROOT / "docs/data-strategy/editions/system-owner/figures"
OUT = HERE / "NASEF2026_IDL_slide_template.pptx"

NAVY, SLATE, STEEL, ASH, MIST, CLOUD = "1E2A44", "4A5878", "7E8AA3", "AEB6C6", "D5DAE3", "EEF0F3"
GREY, INK, WHITE = "5B6272", "161B26", "FFFFFF"
FONT = "Calibri"
FOOT = "Oyewale & Meyer-Petgrave · NASEF 2026 · Petgrave.io"
W, H = 13.333, 7.5
M = 0.6  # side margin

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
count = {"n": 0}


def rgb(h):
    return RGBColor.from_string(h)


def text(slide, x, y, w, h, s, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = line
        f = r.font
        f.name, f.size, f.bold, f.italic, f.color.rgb = FONT, Pt(size), bold, italic, rgb(color)
    return tb


def box(slide, x, y, w, h, fill=None, line=NAVY, dashed=False, radius=True, lw=1.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.06
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dashed:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def placeholder(slide, x, y, w, h, label, detail=""):
    """Dashed box naming what goes here."""
    box(slide, x, y, w, h, fill=WHITE, line=SLATE, dashed=True)
    text(slide, x + 0.2, y + 0.15, w - 0.4, h - 0.3,
         [label] + ([detail] if detail else []), size=14, color=GREY, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


def picture(slide, path, x, y, w=None, h=None):
    if Path(path).exists():
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
    placeholder(slide, x, y, w or 5, h or 3, f"[Figure missing: {Path(path).name}]")


def new_slide(title, section, notes, dark=False):
    count["n"] += 1
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(NAVY if dark else WHITE)
    if not dark:
        text(s, M, 0.25, 3.5, 0.3, section.upper(), size=11, color=STEEL, bold=True)
        text(s, M, 0.55, W - 2 * M, 1.0, title, size=30, color=NAVY, bold=True)
        text(s, M, H - 0.45, 8, 0.3, FOOT, size=10, color=GREY)
        text(s, W - M - 1, H - 0.45, 1, 0.3, str(count["n"]), size=10, color=GREY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def hbars(s, rows, x, y, w, label_w=3.6, row_h=0.75, note=None, focus=0):
    """Native horizontal bars, readable from the back of a room. rows: [(label, value, value_text)]."""
    vmax = max(v for _, v, _ in rows)
    bar_w = w - label_w - 2.4
    for i, (label, v, txt) in enumerate(rows):
        yy = y + i * row_h
        solid = i == focus
        text(s, x, yy, label_w - 0.15, row_h * 0.8, label, size=16, color=INK if solid else GREY, bold=solid,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        bw = max(0.08, bar_w * v / vmax)
        box(s, x + label_w, yy + 0.1, bw, row_h * 0.6, fill=NAVY if solid else ASH, line=None, radius=False)
        text(s, x + label_w + bw + 0.1, yy, 2.3, row_h * 0.8, txt, size=16, color=INK if solid else GREY, bold=solid,
             anchor=MSO_ANCHOR.MIDDLE)
    if note:
        text(s, x, y + len(rows) * row_h + 0.15, w, 0.8, note, size=12, color=GREY)


def card(s, x, y, w, h, big, label, state="solid"):
    fill, line, ink, sub = {
        "solid": (NAVY, NAVY, WHITE, MIST),
        "ash": (CLOUD, ASH, NAVY, GREY),
        "dashed": (WHITE, SLATE, GREY, GREY),
    }[state]
    box(s, x, y, w, h, fill=fill, line=line, dashed=(state == "dashed"))
    text(s, x + 0.25, y + 0.25, w - 0.5, 1.0, big, size=40, color=ink, bold=True)
    text(s, x + 0.25, y + 1.3, w - 0.5, h - 1.45, label, size=15, color=sub)


# 1 · Title
s = new_slide("", "", "20 s. Name, affiliation, one sentence: 'This talk is about turning a solar system's data into answers you can check.'", dark=True)
text(s, M, 1.6, W - 2 * M, 2.2,
     "A reproducible methodology for evidence-grounded conversational diagnostics of solar PV–battery systems",
     size=36, color=WHITE, bold=True)
text(s, M, 3.9, W - 2 * M, 0.5, "Implementation and synthetic evaluation", size=22, color=MIST)
text(s, M, 4.9, 8, 1.0, ["Daniel Oyewola Oyewale · Fitzroy Meyer-Petgrave",
                         "Petgrave.io Technologies Ltd, Nigeria · NASEF 2026"], size=16, color=MIST)
box(s, W - M - 2.0, 4.8, 2.0, 2.0, fill=None, line=ASH, dashed=True)
text(s, W - M - 2.0, 4.8, 2.0, 2.0, ["[QR: live demo]", "engr-daniel.github.io/", "intelligent-data-logger-demo"],
     size=10, color=MIST, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 2 · The logbook
s = new_slide("The system already keeps a diary. Somebody still has to read it", "The problem we noticed",
              "60 s. Tell the radio-station story in your own words: solar, grid and a generator; you read the inverter, meter and "
              "generator hours by hand, wrote them down, and later worked out ROI, solar contribution and consumption for management. "
              "Gaps appeared whenever you were busy. Confirm the details (Q02) and keep the station unnamed.")
placeholder(s, M, 1.8, 5.6, 4.9, "[Photo or drawing]", "A handwritten energy logbook beside an inverter screen")
for i, (h_, b_) in enumerate([("Read by hand", "Inverter, meter and generator hours, whenever someone remembers"),
                              ("Analysed by hand", "ROI, solar contribution and consumption, worked out in the evenings"),
                              ("Acted on late", "Faults found when the power goes off")]):
    y = 1.9 + i * 1.6
    box(s, 6.8, y, W - M - 6.8, 1.35, fill=CLOUD, line=ASH)
    text(s, 7.05, y + 0.12, W - M - 7.3, 0.45, h_, size=20, color=NAVY, bold=True)
    text(s, 7.05, y + 0.6, W - M - 7.3, 0.7, b_, size=15, color=GREY)

# 3 · Scale
s = new_slide("Solar is spreading fast; checking on it is not", "The problem we noticed",
              "60 s. Ember: 1,721 MW imported in the 12 months to June 2025, and no data on how much is installed or working. "
              "Suny installer survey: 6 of 7 had 1–2 return visits a year. The third card needs Q03 (how many of how many installers "
              "reported low maintenance habits). If there is no count, drop the card and say 'installers told us'.")
cw = (W - 2 * M - 0.6) / 3
card(s, M, 2.0, cw, 3.6, "1,721 MW", "of solar panels imported by Nigeria in a year. Installed performance is not tracked. (Ember, 2025)")
card(s, M + cw + 0.3, 2.0, cw, 3.6, "6 of 7", "installers asked had 1–2 jobs a year needing a return visit (Suny survey, 2026)", "ash")
card(s, M + 2 * (cw + 0.3), 2.0, cw, 3.6, "n of N", "installers report low maintenance and check-up habits. [Fill from Q03]", "dashed")

# 4 · Stakes
s = new_slide("Getting the answer wrong costs money, uptime and trust", "Why it matters",
              "60 s. ✂ in a 10-minute slot. The same readings give a 5.1-year or 0.9-year payback depending on whether solar replaced "
              "the grid or a petrol generator. Most calculations assume the grid. Synthetic, 120 days; the petrol figure is an upper bound.")
text(s, M, 1.9, 7.4, 0.5, "Simple payback, same system, same readings", size=16, color=SLATE, bold=True)
hbars(s, [("Solar replaced a petrol generator", 0.9, "0.9 years"),
          ("Solar replaced the grid", 5.1, "5.1 years")], M, 2.6, 7.6, label_w=3.4, row_h=0.95,
      note="₦6.5m system, 120 synthetic days. Petrol 0.8 L/kWh at ₦1,596.25/L (an upper bound); grid ₦225/kWh (demo value). The truth lies between.")
for i, b_ in enumerate(["Money: payback is 5.1 or 0.9 years from the same data, depending on what solar replaced",
                        "Uptime: faults are found when the power goes off",
                        "Trust: one unexplained failure teaches a street that solar is a gamble"]):
    text(s, 8.4, 2.0 + i * 1.4, W - M - 8.4, 1.3, b_, size=16, color=INK if i == 0 else GREY, bold=(i == 0))

# 5 · Solution
s = new_slide("Ask the system a question; get an answer you can check", "The solution",
              "75 s. Describe the outcome before naming the technology: plain-language answers, each with an evidence receipt "
              "(finding, data window, measurements, alternatives, assumptions). Then: 'Claude chooses the tools and writes the "
              "explanation; it never computes the engineering answer.'")
box(s, M, 2.1, 3.4, 1.4, fill=CLOUD, line=ASH)
text(s, M + 0.2, 2.2, 3.0, 1.2, "“Why did our solar output drop on 5 August?”", size=18, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
placeholder(s, M + 3.7, 1.9, 4.4, 4.8, "[Dashboard screenshot]", "The answer, from the live demo")
x0 = M + 8.4
box(s, x0, 1.9, W - M - x0, 4.8, fill=WHITE, line=NAVY)
text(s, x0 + 0.2, 2.0, W - M - x0 - 0.4, 0.4, "Evidence receipt", size=16, color=NAVY, bold=True)
text(s, x0 + 0.2, 2.5, W - M - x0 - 0.4, 4.1,
     ["Finding · candidate cause", "Confidence", "Data window", "Measurements, with units",
      "Alternatives checked", "Data quality", "Assumptions", "Tools used"], size=14, color=GREY)

# 6 · Architecture
s = new_slide("The model chooses tools and explains; engineering truth comes from deterministic analytics", "How it works",
              "75 s. Walk the flow left to right once. Stress the boundary: the model never sees evaluation labels and never computes "
              "an engineering value itself.")
placeholder(s, M, 1.9, W - 2 * M, 4.9, "[Fig2 · End-to-End Evidence-Grounded Workflow]", "Drive › Intelligent Data Logger › Figs")

# 7 · Method
s = new_slide("A method, not just a system: six separations keep the evaluation honest", "The method for building it",
              "75 s. The contribution is the method. Each separation stops a known failure: leaking ground truth, unbounded access, "
              "LLM arithmetic, unverifiable answers, and scoring prose as if it were diagnosis.")
placeholder(s, M, 1.9, 6.6, 4.9, "[Fig1 · Evidence-Grounded Diagnostic Methodology]", "Drive › Figs")
for i, st in enumerate(["Telemetry preparation and validation", "Controlled historical and cross-component access",
                        "Deterministic diagnostic tools", "Structured evidence receipts",
                        "Conversational reporting", "Separate scoring against evaluation-only labels"]):
    y = 1.95 + i * 0.8
    box(s, 7.5, y, 0.55, 0.55, fill=NAVY, line=NAVY)
    text(s, 7.5, y, 0.55, 0.55, str(i + 1), size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.25, y, W - M - 8.25, 0.6, st, size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE)

# 8 · Experiment design
s = new_slide("We tested what each kind of access makes possible", "How we tested it (the demo experiment)",
              "60 s. 2 × 2 design: 60 days vs. 24 hours of history; all components vs. inverter only. Disjoint development and "
              "evaluation installations; frozen protocol; three observation-quality levels.")
placeholder(s, M, 1.9, 6.6, 4.9, "[Fig3 · Information-Access Design]", "Drive › Figs")
cw2 = (W - M - 7.5 - 0.3) / 2
for i, (big, lab) in enumerate([("24", "evaluation installations (8 development, kept apart)"), ("12", "matched episodes each"),
                                ("3", "observation-quality levels"), ("3,456", "paired predictions")]):
    x = 7.5 + (i % 2) * (cw2 + 0.3)
    y = 1.9 + (i // 2) * 2.5
    box(s, x, y, cw2, 2.2, fill=CLOUD, line=ASH)
    text(s, x + 0.2, y + 0.15, cw2 - 0.4, 0.9, big, size=36, color=NAVY, bold=True)
    text(s, x + 0.2, y + 1.1, cw2 - 0.4, 1.0, lab, size=14, color=GREY)

# 9 · Main result
s = new_slide("History across components is what makes diagnosis work", "Results (synthetic)",
              "60 s. Same diagnostic policy throughout; only the information changes. 233 of 240 (95% installation-cluster "
              "interval 95.4–98.8%). An inverter app's last-day view is roughly the bottom bar. These are information "
              "ablations, not competing products.")
hbars(s, [("Full history + all components", 97.1, "97.1%  (233 of 240)"),
          ("Full history + inverter only", 53.8, "53.8%"),
          ("Last 24 h + all components", 40.0, "40.0%"),
          ("Last 24 h + inverter only", 20.0, "20.0%")], M + 0.3, 2.1, W - 2 * M - 0.6, label_w=4.0, row_h=0.95,
      note="Correct diagnoses on identifiable clean cases, same diagnostic policy. Synthetic: 24 evaluation installations, "
           "240 cases per condition; 95% installation-cluster interval for the top bar 95.4–98.8%.")

# 10 · Robustness
s = new_slide("Under heavy corruption it degrades, and it says when it doesn't know", "Results (synthetic)",
              "45 s. ✂ in a 10-minute slot. Severe corruption: full access falls to 90.0%. Restricted conditions often abstain "
              "rather than guess, which is the behaviour we want.")
placeholder(s, M, 1.9, 7.6, 4.9, "[Fig5 · Diagnostic Robustness Under Corruption]", "Drive › Figs")
card(s, 8.6, 1.9, W - M - 8.6, 2.3, "90.0%", "full access under severe corruption")
card(s, 8.6, 4.5, W - M - 8.6, 2.3, "Abstains", "when the evidence is insufficient, instead of guessing", "ash")

# 11 · Conversational pilot
s = new_slide("The language layer picked the right tool every time, and broke the output format every time", "Results (live pilot)",
              "60 s. 24 conversations, one development seed. Every reply was wrapped in a Markdown fence, so the strict JSON contract "
              "failed 24 of 24. With the fences removed, labels matched 24 of 24 but citations were valid in only 15. Fixing this is next.")
placeholder(s, M, 1.9, 6.6, 4.9, "[Fig6 · Conversational Evaluation Failure Decomposition]", "Drive › Figs")
for i, (big, lab, st) in enumerate([("24 / 24", "right tool selected", "solid"), ("0 / 24", "strict output format", "dashed"),
                                     ("24 / 24", "labels, after removing fences", "ash"), ("15 / 24", "citations valid", "ash")]):
    y = 1.9 + i * 1.25
    fill, line = {"solid": (NAVY, NAVY), "ash": (CLOUD, ASH), "dashed": (WHITE, SLATE)}[st]
    box(s, 7.5, y, W - M - 7.5, 1.05, fill=fill, line=line, dashed=(st == "dashed"))
    text(s, 7.7, y + 0.1, 2.2, 0.85, big, size=26, color=WHITE if st == "solid" else NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.9, y + 0.1, W - M - 10.1, 0.85, lab, size=15, color=MIST if st == "solid" else GREY, anchor=MSO_ANCHOR.MIDDLE)

# 12 · Demo
s = new_slide("See it work: three questions", "Demo",
              "90 s. Live if the venue network allows, otherwise the recorded video (Drive: IDL interface Record.mp4). "
              "Q1 shows weather vs. fault. Q2: open the receipt and point at the assumptions. Q3: it refuses to diagnose the data gap.")
placeholder(s, M, 1.9, 8.4, 4.9, "[Live dashboard or video]", "engr-daniel.github.io/intelligent-data-logger-demo/dashboard/")
for i, q in enumerate(["Why did generation drop on 5 August?", "What's our ROI so far?", "What happened during the data gap?"]):
    y = 1.9 + i * 1.7
    box(s, 9.3, y, W - M - 9.3, 1.45, fill=CLOUD, line=ASH)
    text(s, 9.5, y + 0.1, W - M - 9.7, 1.25, [f"Q{i + 1}", q], size=16, color=NAVY, bold=(False), anchor=MSO_ANCHOR.MIDDLE)

# 13 · Limits
s = new_slide("What this shows, and what it does not", "Limits",
              "60 s. Say the limits before the audience does. The contribution is a method, a reference implementation and a "
              "controlled worked evaluation. Field accuracy is the next study.")
half = (W - 2 * M - 0.4) / 2
box(s, M, 1.9, half, 4.9, fill=NAVY, line=NAVY)
text(s, M + 0.3, 2.05, half - 0.6, 0.5, "Shows", size=22, color=WHITE, bold=True)
text(s, M + 0.3, 2.7, half - 0.6, 4.0, ["A reproducible methodology", "A working reference implementation",
                                        "A controlled synthetic evaluation under a frozen protocol",
                                        "Where the language layer fails, measured"], size=17, color=MIST)
box(s, M + half + 0.4, 1.9, half, 4.9, fill=WHITE, line=SLATE, dashed=True)
text(s, M + half + 0.7, 2.05, half - 0.6, 0.5, "Does not show", size=22, color=SLATE, bold=True)
text(s, M + half + 0.7, 2.7, half - 0.6, 4.0, ["Field accuracy: no field data yet",
                                               "Generalisation beyond shared simulator–policy assumptions",
                                               "Live conversational reliability",
                                               "ROI and carbon beyond grid displacement"], size=17, color=GREY)

# 14 · Next
s = new_slide("Next: three real sites, and installers who write down what they find", "What's next",
              "60 s. The pilot needs no new hardware: weekly inverter-counter photos and exports. Installers' notes from return "
              "visits become the reference labels. There is no funding yet. Ask the room for installers and sites.")
placeholder(s, M, 1.9, 6.6, 4.9, "[Fig7 · Field Data Acquisition and Evaluation Architecture]", "Drive › Figs")
box(s, 7.5, 1.9, W - M - 7.5, 4.9, fill=CLOUD, line=ASH)
text(s, 7.75, 2.05, W - M - 8.0, 0.5, "We are looking for", size=20, color=NAVY, bold=True)
text(s, 7.75, 2.7, W - M - 8.0, 3.0, ["Installers willing to log what they find on return visits",
                                       "Two or three sites to share weekly readings for a month",
                                       "Partners interested in field validation"], size=17, color=INK)
text(s, 7.75, 5.9, W - M - 8.0, 0.6, "daniel.oyewale@petgrave.io", size=16, color=NAVY, bold=True)

# 15 · Close
s = new_slide("", "", "20 s. Read the line, then take questions.", dark=True)
text(s, M, 2.3, W - 2 * M, 1.6, "No answer without evidence; no claim without a field check.", size=36, color=WHITE, bold=True)
text(s, M, 4.3, 9, 1.4, ["Demo · engr-daniel.github.io/intelligent-data-logger-demo/dashboard/",
                         "Code · github.com/Engr-Daniel/intelligent-data-logger-demo",
                         "daniel.oyewale@petgrave.io"], size=16, color=MIST)

# Backups
for title, label, notes in [
    ("Backup: why synthetic data?", "[Disjoint seeds · known ground truth · frozen protocol]",
     "No field data or edge device yet. Synthetic installations give known faults, disjoint seeds and a frozen protocol: the methodology claim."),
    ("Backup: ROI and carbon depend on what solar replaced", "",
     "Same 1,851 kWh: 5.1 vs. 0.9 years payback; 883 vs. 3,481 kg CO2. Grid 0.477 kg/kWh (IFI); petrol 2.35 kg/L (EPA). Petrol is an upper bound."),
    ("Backup: how fast does it detect faults?", "[Timing supplement figure]",
     "Secondary, exploratory analysis; see the Timing Supplement."),
    ("Backup: reproduce it", "[README commands · verify_research_run.py]",
     "Fresh clone; run the benchmark; verify an archive without contacting any API."),
]:
    s = new_slide(title, "Backup", notes)
    if "ROI" in title:
        text(s, M, 1.9, 6, 0.5, "Simple payback", size=16, color=SLATE, bold=True)
        hbars(s, [("Petrol generator", 0.9, "0.9 years"), ("Grid", 5.1, "5.1 years")], M, 2.5, 6.0, label_w=2.2, row_h=0.9)
        text(s, M + 6.3, 1.9, 6, 0.5, "CO₂ avoided in 120 days", size=16, color=SLATE, bold=True)
        hbars(s, [("Petrol generator", 3481, "3,481 kg"), ("Grid", 883, "883 kg")], M + 6.3, 2.5, 5.8, label_w=2.2, row_h=0.9)
        text(s, M, 4.7, W - 2 * M, 1.2, "Same 1,851 kWh of solar used on site, synthetic. Grid 0.477 kg CO₂/kWh (IFI harmonised factor); "
             "petrol 2.35 kg CO₂/L (US EPA) at 0.8 L/kWh. The petrol case is an upper bound; the logger can measure the real split "
             "from grid-on and grid-off hours.", size=13, color=GREY)
    else:
        placeholder(s, M, 1.9, W - 2 * M, 4.9, label)

prs.save(OUT)
print("wrote", OUT, "with", count["n"], "slides")
