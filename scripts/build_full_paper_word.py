"""Build current full-paper Word submission, with live TOC/list fields."""
from pathlib import Path
import sys,json,re,argparse
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'.venv/Lib/site-packages'))
from scripts.build_submission import build
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=ROOT/'paper/submission/NASEF2026_Full_Paper_Oyewale_Meyer-Petgrave.docx'
parser=argparse.ArgumentParser()
parser.add_argument('--contact',type=Path,default=ROOT/'.venv/submission-authors.json')
args=parser.parse_args()
contact=json.loads(args.contact.read_text(encoding='utf-8-sig'))
build(ROOT/'paper/FULL_PAPER.md',out,contact)
doc=Document(out);front=json.loads((ROOT/'paper/front_matter.json').read_text())
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(5)
for name in ['Heading 1','Heading 2','Heading 3','Title']:
 style=doc.styles[name];style.paragraph_format.line_spacing=1.08;style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(6);style.font.color.rgb=RGBColor.from_string('203747')
doc.styles['Title'].font.size=Pt(16)
for name in ['Figure Caption','Table Caption','Front Heading']:
 style=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);style.base_style=doc.styles['Normal'];style.font.size=Pt(9 if name!='Front Heading' else 13);style.font.bold=name!='Figure Caption';style.paragraph_format.line_spacing=1;style.paragraph_format.space_after=Pt(7)
doc.styles['Table Caption'].paragraph_format.keep_with_next=True
for p in doc.paragraphs:
 p.paragraph_format.widow_control=True
 if re.match(r'^Figure \d+\.',p.text):p.style='Figure Caption'
 elif re.match(r'^Table \d+\.',p.text):p.style='Table Caption'
 if re.match(r'^\[\d+\]',p.text):p.paragraph_format.keep_together=True
 if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True;p.alignment=1
for table in doc.tables:
 for index,row in enumerate(table.rows):
  for cell in row.cells:
   for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=index<len(table.rows)-1
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
for p in doc.paragraphs[:5]:
 if p.text and not p.text.startswith('Abstract'):p.alignment=1
anchor=next(p for p in doc.paragraphs if p.text=='Abstract')
for text in [front['conference'],front['subtheme']]:
 p=anchor.insert_paragraph_before(text);p.alignment=1;p.paragraph_format.keep_with_next=True
 for r in p.runs:r.bold=True;r.font.size=Pt(10)
intro=next(p for p in doc.paragraphs if p.text=='1. Introduction')
def insert(text='',style=None):
 p=doc.add_paragraph(text,style);intro._p.addprevious(p._p);return p
def page():insert().add_run().add_break(WD_BREAK.PAGE)
def field(code):
 p=insert();r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=code;sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.extend([begin,instr,sep]);p.add_run('Updating contents...');p.add_run()._r.append(end)
page();insert('Table of contents','Front Heading');field('TOC \\o "1-2" \\h \\z \\u')
page();insert('List of figures','Heading 1');field('TOC \\h \\z \\t "Figure Caption,1"')
insert('List of tables','Heading 1');field('TOC \\h \\z \\t "Table Caption,1"')
page();insert('List of abbreviations','Heading 1')
table=doc.add_table(rows=0,cols=2);table.style='Table Grid'
for key,value in front['abbreviations'].items():
 cells=table.add_row().cells;cells[0].text=key;cells[1].text=value
intro._p.addprevious(table._tbl)
page()
setting=OxmlElement('w:updateFields');setting.set(qn('w:val'),'true');doc.settings.element.append(setting)
doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
doc.save(out)
print(out)
