"""Build submission Word files; private contact details stay in a local config."""
from pathlib import Path
import sys,re,json,argparse,hashlib
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.venv/Lib/site-packages'))
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def clean(text):
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',text)
    return text.replace('**','').replace('`','')
def build(source,out,contact):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.9)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=2;normal.paragraph_format.space_after=Pt(0)
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(12);style.font.bold=True
        style.paragraph_format.line_spacing=2;style.paragraph_format.keep_with_next=True
    text=source.read_text(encoding='utf-8-sig');lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            table=doc.add_table(rows=0,cols=len(line.strip('|').split('|')));table.style='Table Grid'
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[clean(c.strip()) for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch('[: -]+',c) for c in cells):
                    row=table.add_row()
                    for cell,value in zip(row.cells,cells):
                        cell.text=value
                        for para in cell.paragraphs:
                            para.paragraph_format.line_spacing=1
                            for run in para.runs:run.font.size=Pt(10)
                    if len(table.rows)==1:
                        props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:tblHeader'))
                        for cell in row.cells:
                            for run in cell.paragraphs[0].runs:run.bold=True
                i+=1
            doc.add_paragraph();continue
        image=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
        if image:
            doc.add_picture(str((source.parent/image.group(2)).resolve()),width=Inches(6.3))
            p=doc.add_paragraph(image.group(1));p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(8)
        elif line.startswith('# '):doc.add_paragraph(line[2:],'Title')
        elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
        elif line.startswith('## '):
            if line=='## Abstract':
                doc.add_paragraph('Corresponding author: '+contact['name']+'; '+contact['email']+'; '+contact['phone'])
            doc.add_paragraph(line[3:],'Heading 1')
        else:
            paragraph=line
            while i+1<len(lines) and lines[i+1].strip() and not lines[i+1].startswith(('#','|','![')):
                i+=1;paragraph+=' '+lines[i].strip()
            p=doc.add_paragraph(clean(paragraph))
            if paragraph.startswith('Table '):p.paragraph_format.keep_with_next=True
        i+=1
    footer=sec.footer.paragraphs[0];footer.alignment=2
    footer.add_run('Page ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title=lines[0].lstrip('# ')
    doc.core_properties.author='Oyewale, D.O.; Meyer-Petgrave, F.'
    doc.core_properties.subject='NASEF 2026: reproducible solar diagnostic methodology and synthetic evaluation'
    doc.save(out)
    return {'file':out.name,'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes)}
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--contact',type=Path,default=ROOT/'.venv/submission-authors.json');args=parser.parse_args()
    contact=json.loads(args.contact.read_text(encoding='utf-8-sig'))
    out=ROOT/'paper/submission';out.mkdir(exist_ok=True)
    records=[build(ROOT/'paper'/source,out/name,contact) for source,name in [('ABSTRACT_DRAFT.md','NASEF2026_Abstract_Oyewale_Meyer-Petgrave.docx'),('FULL_PAPER.md','NASEF2026_Full_Paper_Oyewale_Meyer-Petgrave.docx')]]
    (out/'build_manifest.json').write_text(json.dumps({'format':'A4; 12 pt Times New Roman; double-spaced body; compact tables; provisional full-paper layout','documents':records},indent=2))
    print(json.dumps(records,indent=2))
