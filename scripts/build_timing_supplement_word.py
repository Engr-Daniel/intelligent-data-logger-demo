"""Build the readable secondary timing supplement without rerunning experiments."""
from pathlib import Path
import sys,json
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'.venv/Lib/site-packages'))
from scripts.build_submission import build
from docx import Document
from docx.shared import Pt
out=ROOT/'paper/submission/NASEF2026_Timing_Supplement.docx'
build(ROOT/'paper/TIMING_SUPPLEMENT.md',out,{})
d=Document(out)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Normal'].paragraph_format.space_after=Pt(6)
for name in ['Title','Heading 1','Heading 2']:
 d.styles[name].paragraph_format.line_spacing=1.08
 d.styles[name].paragraph_format.space_after=Pt(8)
d.styles['Title'].font.size=Pt(16)
for p in d.paragraphs:
 if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
 if p.text.startswith('Table S'):p.paragraph_format.keep_with_next=True
for table in d.tables:
 for i,row in enumerate(table.rows):
  for cell in row.cells:
   for p in cell.paragraphs:p.paragraph_format.keep_with_next=i<len(table.rows)-1
d.save(out)
print(out)
